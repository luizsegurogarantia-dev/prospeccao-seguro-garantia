"""Geração do arquivo final consolidado (Excel e/ou CSV).

A carteira traz os dados da licitação, a qualificação da garantia, os links
diretos, os contatos da vencedora e os indicadores de confiança/prioridade.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Config

log = logging.getLogger("seguro_garantia.exporter")

COLUNAS = [
    "Fonte_Origem",
    "Tipo_Registro",
    "Data_Publicacao_Homologacao",
    "Data_Resultado",
    "Estado_UF",
    "Cidade_Orgao",
    "Orgao_Licitante",
    "Objeto_Contratado",
    "Valor_Contrato",
    "Valor_Garantia_Estimado",
    "Link_Edital_Contrato",
    "CNPJ_Vencedor",
    "CNPJ_Valido",
    "Razao_Social_Vencedora",
    "Socios_Administradores",
    "Telefone_Receita",
    "Celular_WhatsApp_Encontrado",
    "Email_Contato",
    "Cidade_Empresa",
    "Setor_Atividade",
    # --- Dossiê da empresa ---
    "Perfil_Empresa",
    "Porte_Empresa",
    "Capital_Social",
    "Data_Abertura",
    "Idade_Empresa",
    "Natureza_Juridica",
    "Simples_Nacional",
    "CNAE_Principal",
    "Sancoes",
    # --- Processo, prazos e vigência ---
    "Numero_Processo",
    "Numero_Edital_Contrato",
    "Data_Assinatura_Contrato",
    "Prazo_Execucao",
    "Vigencia_Contratual",
    # --- Cláusula de garantia do edital/contrato ---
    "Garantia_Status",
    "Garantia_Justificativa",
    "Garantia_Evidencias",
    "Garantia_Percentual_Edital",
    "Prazo_Apresentacao_Garantia",
    "Modalidade_Garantia_Aceita",
    "Vigencia_Garantia",
    # --- Links e rastreabilidade ---
    "Link_Edital",
    "Link_Resultado",
    "Link_Contrato",
    "Link_Ata",
    "Fonte_Dados",
    # --- Qualificação comercial ---
    "Score_Confianca",
    "Prioridade_Comercial",
    "Pendencias",
    # --- Contatos complementares ---
    "Site_Empresa",
    "Outros_Emails",
]

LARGURAS = {
    "Fonte_Origem": 20,
    "Tipo_Registro": 20,
    "Data_Publicacao_Homologacao": 15,
    "Data_Resultado": 15,
    "Estado_UF": 9,
    "Cidade_Orgao": 20,
    "Orgao_Licitante": 38,
    "Objeto_Contratado": 60,
    "Valor_Contrato": 18,
    "Valor_Garantia_Estimado": 18,
    "Link_Edital_Contrato": 32,
    "CNPJ_Vencedor": 20,
    "CNPJ_Valido": 12,
    "Razao_Social_Vencedora": 38,
    "Socios_Administradores": 45,
    "Telefone_Receita": 18,
    "Celular_WhatsApp_Encontrado": 22,
    "Email_Contato": 32,
    "Cidade_Empresa": 26,
    "Setor_Atividade": 28,
    "Perfil_Empresa": 42,
    "Porte_Empresa": 16,
    "Capital_Social": 16,
    "Data_Abertura": 14,
    "Idade_Empresa": 12,
    "Natureza_Juridica": 30,
    "Simples_Nacional": 12,
    "CNAE_Principal": 40,
    "Sancoes": 34,
    "Numero_Processo": 24,
    "Numero_Edital_Contrato": 20,
    "Data_Assinatura_Contrato": 16,
    "Prazo_Execucao": 16,
    "Vigencia_Contratual": 34,
    "Garantia_Status": 18,
    "Garantia_Justificativa": 48,
    "Garantia_Evidencias": 60,
    "Garantia_Percentual_Edital": 14,
    "Prazo_Apresentacao_Garantia": 20,
    "Modalidade_Garantia_Aceita": 26,
    "Vigencia_Garantia": 16,
    "Link_Edital": 32,
    "Link_Resultado": 32,
    "Link_Contrato": 32,
    "Link_Ata": 32,
    "Fonte_Dados": 50,
    "Score_Confianca": 14,
    "Prioridade_Comercial": 18,
    "Pendencias": 50,
    "Site_Empresa": 30,
    "Outros_Emails": 34,
}

COLUNAS_MOEDA = ("Valor_Contrato", "Valor_Garantia_Estimado", "Capital_Social")
COLUNAS_TEXTO_LONGO = (
    "Objeto_Contratado",
    "Socios_Administradores",
    "Sancoes",
    "Perfil_Empresa",
    "Garantia_Justificativa",
    "Garantia_Evidencias",
    "Fonte_Dados",
    "Pendencias",
)
COLUNAS_LINK = (
    "Link_Edital_Contrato",
    "Link_Edital",
    "Link_Resultado",
    "Link_Contrato",
    "Link_Ata",
)

VERDE = "1F6F43"
CINZA_CLARO = "F2F2F2"


def montar_dataframe(linhas: list[dict[str, Any]]) -> pd.DataFrame:
    df = pd.DataFrame(linhas)
    for coluna in COLUNAS:
        if coluna not in df.columns:
            df[coluna] = ""
    df = df[COLUNAS]
    # Ordena pelos resultados mais recentes; empate pelo maior valor de garantia.
    ordem_data = pd.to_datetime(
        df["Data_Resultado"], dayfirst=True, errors="coerce"
    )
    df = (
        df.assign(_ordem_data=ordem_data)
        .sort_values(
            by=["_ordem_data", "Valor_Garantia_Estimado", "Valor_Contrato"],
            ascending=[False, False, False],
            na_position="last",
            kind="stable",
        )
        .drop(columns="_ordem_data")
        .reset_index(drop=True)
    )
    return df


def exportar(df: pd.DataFrame, config: Config) -> dict[str, Path]:
    """Salva o Excel (e opcionalmente o CSV) e devolve os caminhos gerados."""
    config.diretorio_saida.mkdir(parents=True, exist_ok=True)
    carimbo = datetime.now().strftime("%Y%m%d_%H%M")
    base = config.nome_arquivo
    gerados: dict[str, Path] = {}

    caminho_xlsx = config.diretorio_saida / f"{base}_{carimbo}.xlsx"
    _salvar_excel(df, caminho_xlsx, config)
    gerados["xlsx"] = caminho_xlsx

    if config.gerar_csv:
        caminho_csv = config.diretorio_saida / f"{base}_{carimbo}.csv"
        df.to_csv(caminho_csv, sep=";", index=False, encoding="utf-8-sig")
        gerados["csv"] = caminho_csv

    return gerados


def _salvar_excel(df: pd.DataFrame, caminho: Path, config: Config) -> None:
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Prospeccao", index=False)
        planilha = writer.book["Prospeccao"]

        _formatar_cabecalho(planilha)
        _formatar_colunas(planilha, df)
        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        if config.incluir_aba_parametros:
            _aba_parametros(writer.book, config)

    log.info("Excel salvo em %s", caminho)


def _formatar_cabecalho(planilha) -> None:
    preenchimento = PatternFill("solid", fgColor=VERDE)
    fonte = Font(bold=True, color="FFFFFF", size=11)
    for celula in planilha[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    planilha.row_dimensions[1].height = 30


def _formatar_colunas(planilha, df: pd.DataFrame) -> None:
    borda = Border(bottom=Side(style="thin", color="D9D9D9"))
    linhas = len(df) + 1

    for idx, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(idx)
        planilha.column_dimensions[letra].width = LARGURAS.get(coluna, 18)

        for linha in range(2, linhas + 1):
            celula = planilha.cell(row=linha, column=idx)
            celula.border = borda
            if linha % 2 == 0:
                celula.fill = PatternFill("solid", fgColor=CINZA_CLARO)

            if coluna in COLUNAS_MOEDA:
                celula.number_format = '"R$" #,##0.00'
                celula.alignment = Alignment(horizontal="right")
            elif coluna in COLUNAS_TEXTO_LONGO:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
            elif coluna in COLUNAS_LINK and celula.value:
                celula.hyperlink = str(celula.value)
                celula.font = Font(color="0563C1", underline="single")
                celula.alignment = Alignment(vertical="top")
            else:
                celula.alignment = Alignment(vertical="top")

    if "Objeto_Contratado" in df.columns:
        planilha.column_dimensions[
            get_column_letter(list(df.columns).index("Objeto_Contratado") + 1)
        ].width = 60


def _aba_parametros(workbook, config: Config) -> None:
    planilha = workbook.create_sheet("Parametros")
    planilha["A1"] = "Parâmetro"
    planilha["B1"] = "Valor"
    planilha["A1"].font = Font(bold=True)
    planilha["B1"].font = Font(bold=True)
    for i, (chave, valor) in enumerate(config.resumo().items(), start=2):
        planilha.cell(row=i, column=1, value=chave)
        planilha.cell(row=i, column=2, value=str(valor))
    planilha.column_dimensions["A"].width = 32
    planilha.column_dimensions["B"].width = 60
