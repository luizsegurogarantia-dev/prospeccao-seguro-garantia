"""Testes de montagem, ordenação e exportação da carteira."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from seguro_garantia.config import Config
from seguro_garantia.exporter import COLUNAS, exportar, montar_dataframe


def _linha(data: str, valor: float, garantia: float, empresa: str) -> dict:
    return {
        "Data_Resultado": data,
        "Data_Publicacao_Homologacao": data,
        "Valor_Contrato": valor,
        "Valor_Garantia_Estimado": garantia,
        "Razao_Social_Vencedora": empresa,
        "Orgao_Licitante": "Órgão Teste",
        "Garantia_Status": "Provável",
        "Score_Confianca": 70,
        "Prioridade_Comercial": "Média",
        "Link_Edital": "https://pncp.gov.br/app/editais/1",
    }


def test_montar_dataframe_cria_todas_as_colunas():
    df = montar_dataframe([_linha("28/10/2026", 100, 5, "A")])
    assert list(df.columns) == COLUNAS
    assert len(df.columns) >= 50


def test_ordenacao_por_resultado_mais_recente():
    df = montar_dataframe(
        [
            _linha("01/09/2026", 500, 25, "Antiga"),
            _linha("28/10/2026", 100, 5, "Recente"),
            _linha("15/10/2026", 900, 45, "Meio"),
        ]
    )
    assert list(df["Razao_Social_Vencedora"]) == ["Recente", "Meio", "Antiga"]


def test_ordenacao_empate_usa_maior_garantia():
    df = montar_dataframe(
        [
            _linha("28/10/2026", 100, 5, "Menor"),
            _linha("28/10/2026", 900, 45, "Maior"),
        ]
    )
    assert list(df["Razao_Social_Vencedora"]) == ["Maior", "Menor"]


def test_exporta_excel_e_csv(tmp_path):
    df = montar_dataframe(
        [
            _linha("28/10/2026", 100_000, 5_000, "Empresa A"),
            _linha("27/10/2026", 200_000, 10_000, "Empresa B"),
        ]
    )
    config = Config(
        diretorio_saida=tmp_path,
        nome_arquivo="teste",
        gerar_csv=True,
        incluir_aba_parametros=True,
    )

    arquivos = exportar(df, config)

    assert set(arquivos) == {"xlsx", "csv"}
    assert Path(arquivos["xlsx"]).exists()
    assert Path(arquivos["csv"]).exists()

    lido = pd.read_csv(arquivos["csv"], sep=";", encoding="utf-8-sig")
    assert len(lido) == 2
    assert lido.iloc[0]["Razao_Social_Vencedora"] == "Empresa A"
    assert "Garantia_Status" in lido.columns
    assert "Prioridade_Comercial" in lido.columns

    workbook = load_workbook(arquivos["xlsx"])
    assert "Prospeccao" in workbook.sheetnames
    assert "Parametros" in workbook.sheetnames
    planilha = workbook["Prospeccao"]
    cabecalho = [celula.value for celula in planilha[1]]
    coluna_link = cabecalho.index("Link_Edital") + 1
    assert planilha.cell(row=2, column=coluna_link).hyperlink is not None
